"""End-to-end backend tests in MOCK_MODE (no API key needed).

Covers: parsers via API, the stateless analyze->generate contract, edit-mode
semantics (revise in place, delete, insert at position, design preservation,
user overrides), rebuild styling, access-code gate, PDF parsing, error handling.
"""
import json
import os

os.environ["MOCK_MODE"] = "1"
os.environ["UPLOAD_DIR"] = "test_uploads"
os.environ["APP_ACCESS_CODE"] = "test-code-123"  # exercise the access gate
os.environ["APP_USERS"] = '[{"name": "Fatima", "code": "fat-AB12"}]'  # static colleague

from fastapi.testclient import TestClient

import analyzer
import main
import parsers

client = TestClient(main.app)
CODE = "test-code-123"


def ok(cond, msg):
    print(("PASS " if cond else "FAIL ") + msg)
    return bool(cond)


results = []

# ---------- health ----------
r = client.get("/health")
results.append(ok(r.status_code == 200 and r.json()["mock_mode"] and r.json()["access_protected"], f"/health -> {r.json()}"))

# ---------- access gate ----------
with open("samples/old_lesson.pptx", "rb") as fh:
    r = client.post(
        "/analyze",
        data={"objectives": "main idea", "gradeLevel": "3rd Grade", "categories": "Reading", "interactive": "true", "accessCode": "WRONG"},
        files={"file": ("old_lesson.pptx", fh)},
    )
results.append(ok(r.status_code == 401, f"wrong access code rejected with {r.status_code}"))

# ---------- analyze PPTX ----------
with open("samples/old_lesson.pptx", "rb") as fh:
    r = client.post(
        "/analyze",
        data={
            "objectives": "1. Identify the main idea of a text. 2. Use supporting details to justify answers. 3. Retell a story in order.",
            "gradeLevel": "3rd Grade",
            "categories": "Reading",
            "interactive": "true",
            "accessCode": CODE,
        },
        files={"file": ("old_lesson.pptx", fh, "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
results.append(ok(r.status_code == 200, f"/analyze status={r.status_code} {r.text[:200] if r.status_code != 200 else ''}"))
data = r.json()
report = data["report"]
results.append(ok(data["editable_original"] is True, f"editable_original={data['editable_original']} n_units={data['n_units']}"))
results.append(ok(bool(report["slides"]) and bool(report["objectives"]), "report has slides + objectives"))

# ---------- analyze DOCX (worksheet) ----------
with open("samples/old_worksheet.docx", "rb") as fh:
    r2 = client.post(
        "/analyze",
        data={"objectives": "main idea practice", "gradeLevel": "3rd Grade", "categories": "Reading", "interactive": "false", "accessCode": CODE},
        files={"file": ("old_worksheet.docx", fh, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
results.append(ok(r2.status_code == 200 and r2.json()["editable_original"] is False, f"/analyze docx editable_original={r2.json().get('editable_original')}"))

# docx + edit mode must be rejected (file re-sent at generate time is not a pptx)
with open("samples/old_worksheet.docx", "rb") as fh:
    r2gen = client.post(
        "/generate",
        data={"report": __import__("json").dumps(r2.json()["report"]), "mode": "edit", "accessCode": CODE},
        files={"file": ("old_worksheet.docx", fh)},
    )
results.append(ok(r2gen.status_code == 400, f"edit on docx rejected with {r2gen.status_code}"))

# edit mode without any file -> 400
r_nofile = client.post("/generate", data={"report": __import__("json").dumps(report), "mode": "edit", "accessCode": CODE})
results.append(ok(r_nofile.status_code == 400, f"edit without file rejected with {r_nofile.status_code}"))

# ---------- generate: EDIT mode, AI defaults (file re-sent, stateless) ----------
with open("samples/old_lesson.pptx", "rb") as fh:
    r3 = client.post(
        "/generate",
        data={"report": __import__("json").dumps(report), "mode": "edit", "slideActions": "{}", "includeNewSlides": "true", "accessCode": CODE},
        files={"file": ("old_lesson.pptx", fh)},
    )
results.append(ok(r3.status_code == 200, f"/generate edit status={r3.status_code} {r3.text[:200] if r3.status_code != 200 else ''}"))
out1 = "samples/out_edit_default.pptx"
with open(out1, "wb") as fh:
    fh.write(r3.content)

from pptx import Presentation

prs = Presentation(out1)
texts = [(s.shapes.title.text if s.shapes.title is not None else "") for s in prs.slides]
print("   edit-default titles:", texts)
# mock report: u1 keep, u2 revise, u3 remove, u4 revise, u5 remove, u6 revise, u7 n/a
# expected: survivors [1,2,4,6,7] + new slide after unit 1 => 6 slides, new slide at position 2
results.append(ok(len(prs.slides) == 6, f"slide count edit-default = {len(prs.slides)} (expect 6)"))
results.append(ok("Revised slide 2" in texts[2], f"revised unit-2 slide present at pos 3: {texts[2]!r}"))
results.append(ok("Supporting details" in texts[1], f"new gap slide inserted at position 2: {texts[1]!r}"))
results.append(ok(not any("Silent E" in t for t in texts), "removed slides (3,5) gone"))
prsslides = list(prs.slides)
results.append(ok(prsslides[2].slide_layout.name == "Title and Content", f"design kept on revised slide (layout={prsslides[2].slide_layout.name})"))

# ---------- generate: EDIT mode with user overrides (flat map accepted) ----------
with open("samples/old_lesson.pptx", "rb") as fh:
    r4 = client.post(
        "/generate",
        data={"report": __import__("json").dumps(report), "mode": "edit", "slideActions": '{"3": "keep", "6": "keep"}', "includeNewSlides": "true", "accessCode": CODE},
        files={"file": ("old_lesson.pptx", fh)},
    )
out2 = "samples/out_edit_override.pptx"
with open(out2, "wb") as fh:
    fh.write(r4.content)
prs2 = Presentation(out2)
texts2 = [(s.shapes.title.text if s.shapes.title is not None else "") for s in prs2.slides]
print("   edit-override titles:", texts2)
results.append(ok(len(prs2.slides) == 7, f"slide count edit-override = {len(prs2.slides)} (expect 7: one removal + 1 new)"))
results.append(ok(any("Example: The Lost Dog" in t for t in texts2), "user override kept slide 3"))

# ---------- generate: REBUILD mode (no file needed) ----------
r5 = client.post("/generate", data={"report": __import__("json").dumps(report), "mode": "rebuild", "gradeLevel": "3rd Grade", "categories": "Reading", "sourceName": "old_lesson.pptx", "accessCode": CODE})
results.append(ok(r5.status_code == 200, f"/generate rebuild status={r5.status_code}"))
out3 = "samples/out_rebuild.pptx"
with open(out3, "wb") as fh:
    fh.write(r5.content)
prs3 = Presentation(out3)
first_content = prs3.slides[1]
font_ok = False
for shape in first_content.shapes:
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                if run.font.name == "Comic Sans MS" and run.font.size and run.font.size.pt == 36:
                    font_ok = True
results.append(ok(font_ok, "rebuild deck styled Comic Sans 36pt"))
print("   rebuild titles:", [(s.shapes.title.text if s.shapes.title is not None else "") for s in prs3.slides])

# ---------- generate: wrong access code ----------
r6 = client.post("/generate", data={"report": __import__("json").dumps(report), "mode": "rebuild", "accessCode": "nope"})
results.append(ok(r6.status_code == 401, f"generate with wrong code -> {r6.status_code}"))

# ---------- generate: corrupt report JSON ----------
r7 = client.post("/generate", data={"report": "{not json", "mode": "rebuild", "accessCode": CODE})
results.append(ok(r7.status_code == 400, f"corrupt report JSON -> {r7.status_code}"))

# ---------- multi-user: login + colleague code ----------
r_l = client.get("/login?code=fat-AB12")
results.append(ok(r_l.status_code == 200 and r_l.json().get("name") == "Fatima", f"/login colleague -> {r_l.status_code}"))
r_l2 = client.get("/login?code=test-code-123")
results.append(ok(r_l2.status_code == 200 and r_l2.json().get("role") == "admin", "/login admin role ok"))
r_l3 = client.get("/login?code=bogus")
results.append(ok(r_l3.status_code == 401, f"/login bad code -> {r_l3.status_code}"))
r_an = client.post(
    "/analyze",
    data={"objectives": "x", "gradeLevel": "1st Grade", "categories": "Reading", "interactive": "false", "accessCode": "fat-AB12"},
)
results.append(ok(r_an.status_code == 400, f"colleague code passes auth gate (400=missing file) -> {r_an.status_code}"))

# ---------- design mode (no report, no file) ----------
r_d = client.post(
    "/generate",
    data={"designMode": "true", "objectives": "1. main idea 2. retelling", "mode": "rebuild",
          "gradeLevel": "3rd Grade", "categories": "Reading", "interactive": "true",
          "userName": "Fatima", "accessCode": CODE},
)
results.append(ok(r_d.status_code == 200, f"/generate designMode -> {r_d.status_code}"))
out4 = "samples/out_design.pptx"
with open(out4, "wb") as fh:
    fh.write(r_d.content)
from pptx import Presentation as _P
prs4 = _P(out4)
t4 = [(s.shapes.title.text if s.shapes.title is not None else "") for s in prs4.slides]
results.append(ok(len(prs4.slides) >= 3 and any("Warm-up" in t for t in t4), f"design deck: {len(prs4.slides)} slides, titles={t4[:3]}"))

# ---------- parse_pdf on a hand-crafted minimal PDF ----------
def make_minimal_pdf(path, line):
    objs = []
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
    objs.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>")
    stream = f"BT /F1 24 Tf 72 720 Td ({line}) Tj ET".encode()
    objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj ".encode() + body + b" endobj\n"
    xref_pos = len(out)
    out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += b"trailer\n<< /Size " + str(len(objs) + 1).encode() + b" /Root 1 0 R >>\nstartxref\n" + str(xref_pos).encode() + b"\n%%EOF"
    with open(path, "wb") as fh:
        fh.write(out)

make_minimal_pdf("samples/old_slides.pdf", "Main Idea Worksheet - Page One")
mp = parsers.parse_pdf("samples/old_slides.pdf")
results.append(ok(mp.kind == "pdf" and "Main Idea" in mp.units[0].title, f"parse_pdf -> {len(mp.units)} unit, title={mp.units[0].title!r}"))

# ---------- /illustrate (mock) ----------
r = client.post("/illustrate", data={"report": json.dumps(analyzer.mock_report(4)), "accessCode": CODE})
body = r.json() if r.status_code == 200 else {}
results.append(ok(
    r.status_code == 200 and body.get("mock") and len(body.get("illustrations", [])) >= 1
    and body["illustrations"][0]["image_url"] == "",
    f"/illustrate mock -> {r.status_code}, {len(body.get('illustrations', []))} placeholder(s)",
))
r = client.post("/illustrate", data={"report": "not json", "accessCode": CODE})
results.append(ok(r.status_code == 400, f"/illustrate bad report -> {r.status_code}"))

# ---------- /generate embeds illustrations (graceful on dead URLs) ----------
report = analyzer.mock_report(4)
illo_title = report["new_slides"][0]["title"]
fd = {
    "report": json.dumps(report),
    "mode": "rebuild",
    "accessCode": CODE,
    "gradeLevel": "3rd Grade",
    "illustrations": json.dumps([
        {"title": illo_title, "url": ""},                                        # empty -> ignored
        {"title": "no such slide", "url": "https://127.0.0.1:1/dead.png"},      # dead URL -> skipped
    ]),
}
r = client.post("/generate", data=fd)
results.append(ok(r.status_code == 200 and r.headers["content-type"].startswith("application/vnd"), f"/generate with illustrations (dead URL tolerated) -> {r.status_code}"))
if r.status_code == 200:
    out = os.path.join("test_uploads", "illo_rebuild.pptx")
    open(out, "wb").write(r.content)
    from pptx import Presentation
    prs = Presentation(out)
    slide_titles = [s.shapes.title.text for s in prs.slides if s.shapes.title is not None]
    results.append(ok(illo_title in slide_titles, f"new slide still present: {illo_title!r} in {slide_titles[:6]}"))
    pic_slides = [i for i, s in enumerate(prs.slides) for sh in s.shapes if sh.shape_type == 13]  # PICTURE
    results.append(ok(len(pic_slides) == 0, f"no picture shapes from dead URLs: {pic_slides}"))

print(f"\n{'='*50}\n{sum(results)}/{len(results)} checks passed")
raise SystemExit(0 if all(results) else 1)
