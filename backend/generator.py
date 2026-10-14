"""Build the updated lesson PPTX.

Two modes:
  rebuild_deck() – generate a fresh deck (teacher's original style: Comic Sans, 36pt,
                   max a few bullets per slide) from the analysis report + context.
  edit_deck()    – take the ORIGINAL uploaded .pptx and apply the report:
                     verdict "revise" -> rewrite title/body text in place (design kept)
                     verdict "remove" -> delete the slide
                     report.new_slides -> insert new slides at the requested positions,
                                          using the source deck's own layouts.
"""
from __future__ import annotations

import os

from pptx import Presentation
from pptx.util import Pt

REBUILD_FONT = os.getenv("DECK_FONT", "Comic Sans MS")
REBUILD_FONT_SIZE = int(os.getenv("DECK_FONT_SIZE", "36"))
MAX_BULLETS = 5

_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


# --------------------------------------------------------------------------- #
# Shared helpers
# --------------------------------------------------------------------------- #
def _content_lines(content: str) -> list[str]:
    lines = [ln.strip("•-– ").strip() for ln in (content or "").splitlines()]
    return [ln for ln in lines if ln][:MAX_BULLETS]


def _style_runs(slide, font_name: str, size: int):
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.name = font_name
                run.font.size = Pt(size)


def _title_placeholder(slide):
    try:
        if slide.shapes.title is not None:
            return slide.shapes.title
    except Exception:
        pass
    for ph in slide.placeholders:
        if ph.has_text_frame and ph.placeholder_format.idx == 0:
            return ph
    return None


def _body_placeholder(slide):
    for ph in slide.placeholders:
        if ph.has_text_frame and ph.placeholder_format.idx != 0:
            return ph
    return None


def _set_title(slide, title: str) -> bool:
    ph = _title_placeholder(slide)
    if ph is None:
        return False
    ph.text_frame.text = title or ""
    return True


def _set_body(slide, lines: list[str]) -> bool:
    ph = _body_placeholder(slide)
    if ph is None:
        return False
    tf = ph.text_frame
    tf.clear()
    if lines:
        tf.text = lines[0]
        for line in lines[1:]:
            p = tf.add_paragraph()
            p.text = line
    return True


def _delete_slide(prs, zero_based_index: int) -> None:
    """Remove a slide from the deck (python-pptx has no public API for this)."""
    sld_id_lst = prs.slides._sldIdLst
    sld_id = list(sld_id_lst)[zero_based_index]
    prs.part.drop_rel(sld_id.get(_R_ID))
    sld_id_lst.remove(sld_id)


def _pick_layout(prs, near_slide=None):
    """Prefer the layout of a neighbouring slide; else a titled bullet layout."""
    if near_slide is not None:
        return near_slide.slide_layout
    for layout in prs.slide_layouts:
        has_title = any(ph.placeholder_format.idx == 0 for ph in layout.placeholders)
        has_body = any(ph.placeholder_format.idx != 0 for ph in layout.placeholders)
        if has_title and has_body:
            return layout
    return prs.slide_layouts[min(1, len(prs.slide_layouts) - 1)]


def _final_action(entry: dict, overrides: dict) -> str:
    return overrides.get(str(entry.get("index")), entry.get("verdict", "keep"))


# --------------------------------------------------------------------------- #
# Illustrations (optional images embedded into new slides)
# --------------------------------------------------------------------------- #
def _illustration_map(entries) -> dict[str, str]:
    """[{"title","url"}] -> {title: url}, later duplicates win."""
    return {str(e["title"]): str(e["url"]) for e in (entries or []) if e.get("title") and e.get("url")}


def _fetch_image(url: str, dest_dir: str) -> str | None:
    """Download an illustration; None on any failure (deck build must never die for art)."""
    import uuid as _uuid

    import httpx

    try:
        r = httpx.get(url, timeout=20.0, follow_redirects=True)
        r.raise_for_status()
        if len(r.content) < 1000:  # placeholder/garbage guard
            return None
        path = os.path.join(dest_dir, f"illo_{_uuid.uuid4().hex[:8]}")
        open(path, "wb").write(r.content)
        return path
    except Exception as exc:
        print(f"[generator] illustration download failed, skipping: {exc.__class__.__name__}: {str(exc)[:120]}")
        return None


def _add_illustration(prs, slide, img_path: str) -> None:
    """Place the image on the right side of the slide, narrowing the body text."""
    from pptx.util import Inches

    try:
        body = _body_placeholder(slide)
        if body is not None and body.width and body.width > prs.slide_width * 0.55:
            body.width = int(prs.slide_width * 0.52)
        img_w = int(prs.slide_width * 0.40)
        left = prs.slide_width - img_w - Inches(0.35)
        top = int(prs.slide_height * 0.32)
        slide.shapes.add_picture(img_path, left, top, width=img_w)
    except Exception as exc:
        print(f"[generator] illustration embed failed, skipping: {exc.__class__.__name__}: {str(exc)[:120]}")


# --------------------------------------------------------------------------- #
# Mode A: rebuild fresh deck
# --------------------------------------------------------------------------- #
def rebuild_deck(report: dict, meta: dict, out_path: str) -> str:
    """meta: {grade_level, categories, interactive, source_name, illustrations?}"""
    prs = Presentation()
    illo_by_title = _illustration_map(meta.get("illustrations"))

    title_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_layout)
    slide.shapes.title.text = _deck_title(report)
    if len(slide.placeholders) > 1 and slide.placeholders[1].has_text_frame:
        slide.placeholders[1].text_frame.text = (
            f"{meta.get('grade_level', '')} • {meta.get('categories', '')}\n"
            f"Updated from: {meta.get('source_name', 'previous material')}"
        ).strip()
    _style_runs(slide, REBUILD_FONT, REBUILD_FONT_SIZE)

    body_layout = prs.slide_layouts[min(1, len(prs.slide_layouts) - 1)]

    def add_body_slide(slide_title: str, content: str):
        s = prs.slides.add_slide(body_layout)
        _set_title(s, slide_title or "")
        _set_body(s, _content_lines(content))
        _style_runs(s, REBUILD_FONT, REBUILD_FONT_SIZE)
        url = illo_by_title.get(slide_title or "")
        if url:
            img = _fetch_image(url, os.path.dirname(out_path))
            if img:
                _add_illustration(prs, s, img)

    for sl in report.get("slides", []):
        if sl.get("verdict") == "remove":
            continue
        if sl.get("verdict") == "revise" and sl.get("revised_content"):
            add_body_slide(sl.get("revised_title") or sl.get("title", ""), sl["revised_content"])
        elif sl.get("content"):
            add_body_slide(sl.get("title", ""), sl["content"])
        else:
            add_body_slide(sl.get("title", ""), "(content carried over from previous year)")
    for ns in report.get("new_slides", []):
        add_body_slide(ns.get("title", ""), ns.get("content", ""))

    prs.save(out_path)
    return out_path


def _deck_title(report: dict) -> str:
    objectives = report.get("objectives") or []
    if objectives:
        return f"Lesson: {objectives[0].get('text', '')}"[:120]
    return "Updated Lesson Presentation"


# --------------------------------------------------------------------------- #
# Mode B: edit the original deck, preserving its design
# --------------------------------------------------------------------------- #
def edit_deck(source_path: str, report: dict, actions: dict, out_path: str) -> str:
    """actions: {
         "slide_actions": {"<unit index>": "keep|revise|remove"},  # optional override
         "include_new_slides": bool  (default True)
       }
    Unit indexes refer to the ORIGINAL uploaded deck (1-based).
    """
    import shutil

    shutil.copyfile(source_path, out_path)
    prs = Presentation(out_path)

    # Accept either {"slide_actions": {...}, "include_new_slides": bool}
    # or a flat override map {"<unit index>": "keep|revise|remove"}.
    if "slide_actions" in actions:
        overrides: dict = actions.get("slide_actions") or {}
    else:
        overrides = {k: v for k, v in actions.items() if k != "include_new_slides"}
    include_new: bool = actions.get("include_new_slides", True)

    slide_entries = {int(s["index"]): s for s in report.get("slides", [])}
    unit_count = len(list(prs.slides))

    # ---- 1. in-place revisions (original design untouched) ---------------
    revised = []
    for idx, entry in slide_entries.items():
        if not 1 <= idx <= unit_count or _final_action(entry, overrides) != "revise":
            continue
        if not (entry.get("revised_title") or entry.get("revised_content")):
            continue
        slide = list(prs.slides)[idx - 1]
        changed = False
        if entry.get("revised_title"):
            changed |= _set_title(slide, entry["revised_title"])
        if entry.get("revised_content"):
            changed |= _set_body(slide, _content_lines(entry["revised_content"]))
        if changed:
            revised.append(idx)

    # ---- 2. plan: what gets removed / kept / inserted ---------------------
    to_remove = [
        idx
        for idx, entry in slide_entries.items()
        if 1 <= idx <= unit_count and _final_action(entry, overrides) == "remove"
    ]
    survivors = [i for i in range(1, unit_count + 1) if i not in set(to_remove)]
    new_specs = [ns for ns in report.get("new_slides", []) if ns.get("title")] if include_new else []

    sld_id_lst = prs.slides._sldIdLst
    orig_id_by_unit = {unit: sld_id for unit, sld_id in enumerate(sld_id_lst, start=1)}

    # ---- 3. append new slides FIRST, while all original parts still exist.
    # python-pptx names a new slide part "slide{len(slideList)+1}.xml", so
    # deleting before adding makes that name collide with a surviving part and
    # corrupts the saved zip. Appending first + a uniqueness sweep is safe.
    inserted = 0
    new_id_by_spec: dict = {}
    illo_by_title = _illustration_map(actions.get("illustrations"))
    for spec in new_specs:
        after = int(spec.get("after_index") or 0)
        neighbour = max((i for i in survivors if i <= after), default=None)
        layout_source = None
        if neighbour is not None:
            layout_source = list(prs.slides)[survivors.index(neighbour)]
        s = prs.slides.add_slide(_pick_layout(prs, layout_source))
        _set_title(s, spec.get("title", ""))
        _set_body(s, _content_lines(spec.get("content", "")))
        url = illo_by_title.get(spec.get("title", ""))
        if url:
            img = _fetch_image(url, os.path.dirname(out_path))
            if img:
                _add_illustration(prs, s, img)
        new_id_by_spec[id(spec)] = list(sld_id_lst)[-1]
        inserted += 1

    # ---- 4. deletions (highest index first so indexes stay valid) ---------
    for idx in sorted(to_remove, reverse=True):
        _delete_slide(prs, idx - 1)

    # ---- 5. one reorder pass (appending an existing element MOVES it) -----
    if new_id_by_spec:
        order: list[tuple] = [("orig", i) for i in survivors]
        for spec in new_specs:
            after = int(spec.get("after_index") or 0)
            pos = 0
            for j, token in enumerate(order):
                anchor = token[1] if token[0] == "orig" else int(token[1].get("after_index") or 0)
                if anchor <= after:
                    pos = j + 1
            order.insert(pos, ("new", spec))
        for token_type, payload in order:
            sld_id = orig_id_by_unit[payload] if token_type == "orig" else new_id_by_spec[id(payload)]
            sld_id_lst.append(sld_id)

    _ensure_unique_slide_partnames(prs)
    prs.save(out_path)
    print(f"[edit_deck] revised={sorted(revised)} removed={sorted(to_remove, reverse=True)} inserted={inserted}")
    return out_path


def _ensure_unique_slide_partnames(prs) -> None:
    """Guarantee no two slide parts share a partname before serialization."""
    from pptx.opc.packuri import PackURI

    used: set[str] = set()
    for sld_id in prs.slides._sldIdLst:
        part = prs.part.related_part(sld_id.get(_R_ID))
        if str(part.partname) in used:
            n = 1
            while f"/ppt/slides/slide{n}.xml" in used:
                n += 1
            part.partname = PackURI(f"/ppt/slides/slide{n}.xml")
        used.add(str(part.partname))
